#!/usr/bin/env python3
"""CSL Plasma megrendelő + teljesítésigazolás a továbbszámlázott tételekről.

A CSL saját formáját használjuk, ezért nem építjük újra a dokumentumokat, hanem az
assets/ alatti két sablon másolatát töltjük ki — így a fejléc, a betűk és a
táblázatstílus változatlan marad.

A tételsorokat MINDIG a végleges számláról vedd át, ne a munkatáblából: a kettő
el szokott térni (összevont sorok, kerekített egységár), és a három dokumentumnak
forintra egyeznie kell.

    python3 scripts/csl_megrendelo_tig.py --minta > csl.json
    python3 scripts/csl_megrendelo_tig.py csl.json --kimenet ./out
"""
import argparse, copy, json, sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

NBSP = ' '
GYOKER = Path(__file__).resolve().parent.parent
MEGR_SABLON = GYOKER / 'assets' / 'csl_megrendelo_sablon.docx'
TIG_SABLON = GYOKER / 'assets' / 'csl_tig_sablon.docx'
ROMAI = ['I.', 'II.', 'III.', 'IV.', 'V.', 'VI.', 'VII.', 'VIII.']

MINTA = {
    "megrendeles_helye": "Szeged",
    "megrendeles_datum": "2026.08.11.",
    "tig_helye": "Budapest",
    "tig_datum": "2026.09.11.",
    "fizetesi_hatarido": "2026.09.19.",
    "megrendeles_targya": "Nyomdai anyagok, szoftverek beszerzése, szállítása",
    "szakaszok": [
        {
            "fejezet": None,
            "cim": "CSL Club – szoftverek, nyomdai anyagok",
            "tig_cim": "CSL Club – szoftver és nyomda",
            "po": "4400911742",
            "tetelek": [
                ["Plasztikkártya", 600, "db", 82110],
                ["Szállítás", 1, "db", 2490],
            ],
        },
        {
            "fejezet": None,
            "cim": "Egyéb szoftverek, nyomdai anyagok",
            "tig_cim": "Szoftver és nyomda",
            "po": "4400911740",
            "tetelek": [
                ["Plazma kérdőív", 18000, "db", 88200],
                ["Reklámhordozó EPR", 44.76, "kg", 4207],
            ],
        },
    ],
}


def ezres(n):
    return f'{int(round(n)):,}'.replace(',', NBSP)


def mennyiseg(n):
    if abs(n - round(n)) < 1e-9:
        return ezres(n)
    s = f'{n:,.2f}'.replace(',', NBSP).replace('.', ',')
    return s.rstrip('0').rstrip(',')


def settext(p, txt):
    """Az első run formázását megtartva cseréli a bekezdés szövegét."""
    if not p.runs:
        p.add_run(txt)
        return
    p.runs[0].text = txt
    for r in p.runs[1:]:
        r.text = ''


def cellset(cell, txt):
    settext(cell.paragraphs[0], txt)
    for p in cell.paragraphs[1:]:
        for r in p.runs:
            r.text = ''


def sorszam(tbl, n):
    """A tábla sorainak számát az utolsó sor klónozásával/törlésével állítja be."""
    while len(tbl.rows) > n:
        tbl._tbl.remove(tbl.rows[-1]._tr)
    while len(tbl.rows) < n:
        tbl._tbl.append(copy.deepcopy(tbl.rows[-1]._tr))


def blokkok(doc):
    """A body gyermekei sorban, (tipus, elem, wrapper) hármasokként."""
    ki = []
    for ch in doc.element.body.iterchildren():
        tag = ch.tag.split('}')[1]
        if tag == 'p':
            ki.append(('p', ch, Paragraph(ch, doc)))
        elif tag == 'tbl':
            ki.append(('tbl', ch, Table(ch, doc)))
        else:
            ki.append((tag, ch, None))
    return ki


def megrendelo(adat, kimenet):
    doc = Document(str(MEGR_SABLON))
    b = blokkok(doc)
    p_idx = [i for i, (t, _, _) in enumerate(b) if t == 'p']

    def p(n):  # n-edik bekezdés
        return b[p_idx[n]]

    # sablonblokkok kimentése az első szakaszból
    minta_fejezet = copy.deepcopy(p(8)[1])      # "I. Debrecen – Campus Fesztivál"
    minta_cim = copy.deepcopy(p(10)[1])         # "… (PO: …):"
    minta_ures = copy.deepcopy(p(9)[1])
    minta_tbl = copy.deepcopy([e for t, e, _ in b if t == 'tbl'][0])
    minta_dij = copy.deepcopy(p(13)[1])         # "A fentiek díja: … Ft + ÁFA"

    # a régi szakaszok (P8 … P44 és minden tábla) törlése
    eleje, vege = p_idx[8], p_idx[44]
    horgony = b[p_idx[45]][1]
    for t, el, _ in b[eleje:vege + 1]:
        el.getparent().remove(el)

    ossz = 0
    fejezet_db = 0
    for sz in adat['szakaszok']:
        tetelek = sz['tetelek']
        reszossz = sum(x[3] for x in tetelek)
        ossz += reszossz
        if sz.get('fejezet'):
            fejezet_db += 1
            e = copy.deepcopy(minta_fejezet)
            horgony.addprevious(e)
            settext(Paragraph(e, doc), f"{ROMAI[fejezet_db - 1]} {sz['fejezet']}")
            horgony.addprevious(copy.deepcopy(minta_ures))
        e = copy.deepcopy(minta_cim)
        horgony.addprevious(e)
        settext(Paragraph(e, doc), f"{sz['cim']} (PO: {sz['po']}):")
        horgony.addprevious(copy.deepcopy(minta_ures))

        e = copy.deepcopy(minta_tbl)
        horgony.addprevious(e)
        tbl = Table(e, doc)
        sorszam(tbl, len(tetelek))
        for r, (megn, menny, egys, netto) in zip(tbl.rows, tetelek):
            cellset(r.cells[0], str(megn))
            cellset(r.cells[1], mennyiseg(menny))
            cellset(r.cells[2], str(egys))
            cellset(r.cells[3], ezres(netto) + ' Ft')

        horgony.addprevious(copy.deepcopy(minta_ures))
        e = copy.deepcopy(minta_dij)
        horgony.addprevious(e)
        settext(Paragraph(e, doc), f'A fentiek díja: {ezres(reszossz)} Ft + ÁFA')
        horgony.addprevious(copy.deepcopy(minta_ures))

    b2 = blokkok(doc)
    p2 = [i for i, (t, _, _) in enumerate(b2) if t == 'p']
    for i in p2:
        szoveg = b2[i][2].text
        if szoveg.startswith('Mindösszesen:'):
            settext(b2[i][2], f'Mindösszesen: {ezres(ossz)} Ft + ÁFA')
        elif szoveg.startswith('Szeged,') or szoveg.startswith('Budapest,'):
            settext(b2[i][2], f"{adat.get('megrendeles_helye', 'Szeged')}, {adat['megrendeles_datum']}")
    doc.save(str(kimenet))
    return ossz


def tig(adat, kimenet, ossz):
    doc = Document(str(TIG_SABLON))
    b = blokkok(doc)
    tablak = [w for t, _, w in b if t == 'tbl']
    fej, osszeg = tablak[0], tablak[1]

    cellset(fej.rows[0].cells[1], ' ' + adat['megrendeles_datum'])
    cellset(fej.rows[1].cells[1], adat.get('szerzodo_partner', 'Marketing Store Kft.'))
    cellset(fej.rows[2].cells[1], adat['megrendeles_targya'])

    # az utolsó két sor (fizetési mód / határidő) formázása eltér — megőrizzük
    farok = [copy.deepcopy(osszeg.rows[-2]._tr), copy.deepcopy(osszeg.rows[-1]._tr)]
    sorszam(osszeg, 2 + len(adat['szakaszok']))  # "Összege:" + szakaszok + "Mindösszesen"
    for tr in farok:
        osszeg._tbl.append(tr)

    sorok = [('Összege:', '')]
    for sz in adat['szakaszok']:
        reszossz = sum(x[3] for x in sz['tetelek'])
        cim = sz.get('tig_cim', sz['cim'])
        sorok.append((f"{cim} (PO {sz['po']})", f'{ezres(reszossz)} + (27% áfa)'))
    sorok.append(('Mindösszesen', f'{ezres(ossz)} + (27% áfa)'))
    sorok.append(('Fizetés módja: ', adat.get('fizetes_modja', 'Átutalás')))
    sorok.append(('Fizetés határideje:', ' ' + adat['fizetesi_hatarido']))

    for r, (bal, jobb) in zip(osszeg.rows, sorok):
        cellset(r.cells[0], bal)
        cellset(r.cells[1], jobb)

    for t, _, w in b:
        if t == 'p' and (w.text.startswith('Budapest,') or w.text.startswith('Szeged,')):
            settext(w, f"{adat.get('tig_helye', 'Budapest')}, {adat['tig_datum']}")
    doc.save(str(kimenet))


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('json', nargs='?', help='bemeneti JSON')
    ap.add_argument('--minta', action='store_true', help='mintafájl kiírása')
    ap.add_argument('--kimenet', default='.', help='kimeneti könyvtár')
    a = ap.parse_args()

    if a.minta:
        json.dump(MINTA, sys.stdout, ensure_ascii=False, indent=2)
        print()
        return
    if not a.json:
        ap.error('adj meg egy JSON fájlt, vagy használd a --minta kapcsolót')

    adat = json.loads(Path(a.json).read_text(encoding='utf-8'))
    ki = Path(a.kimenet)
    ki.mkdir(parents=True, exist_ok=True)

    mnev = ki / f"Megrendelés {adat['megrendeles_datum']} CSL Plasma, Kampány szoftverek.docx"
    tnev = ki / f"Teljesítés igazolás {adat['tig_datum']} CSL Plasma, Kampány szoftverek.docx"
    ossz = megrendelo(adat, mnev)
    tig(adat, tnev, ossz)

    print(f'Megrendelő : {mnev}')
    print(f'TIG        : {tnev}')
    for sz in adat['szakaszok']:
        print(f"  {sz['cim']} (PO {sz['po']}): {ezres(sum(x[3] for x in sz['tetelek']))} Ft")
    print(f'  MINDÖSSZESEN: {ezres(ossz)} Ft + ÁFA')


if __name__ == '__main__':
    main()
